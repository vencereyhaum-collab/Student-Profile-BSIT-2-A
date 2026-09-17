from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

base_dir = Path(__file__).resolve().parent
output_path = base_dir / 'Haum Resume.docx'
image_path = base_dir / 'ID Picture.jpg'

profile_data = {
    'name': 'Vence Rey Haum',
    'role': 'Bachelor of Science in Information Technology – 1st year',
    'address': 'Baybay-1 Poblacion Bislig City',
    'email': 'vencereyhaum@gmail.com',
    'mobile': '09519693197',
    'facebook': 'haum.rey',
    'summary': 'I am an Information Technology student who enjoys learning about programming, responsible, and always eager to improve my skills.',
    'career_goals': 'To become a successful IT professional and develop systems that help people and businesses.',
    'education': [
        'Elementary School: Bislig Central Elementary School',
        'High School: Bislig City National High School',
        'Senior High School: Bislig City National High School',
        'College (Current): De La Salle John Bosco College',
    ],
    'technical_skills': ['Basic Programming'],
    'programming_languages': ['Java'],
    'software_applications': ['Canva'],
    'soft_skills': ['Communication', 'Teamwork', 'Time Management'],
    'project_title': 'Student Information System',
    'project_description': 'A web-based system for managing student records efficiently.',
    'project_language': 'Java',
    'project_status': 'Completed',
    'project_link': 'https://example.com/student-information-system',
    'achievements': [
        'Certificates Earned: Certificate of Participation',
        'Academic Awards: None',
        'Seminars and Workshops: Not specified',
        'Work Experience: None',
    ],
    'motto': 'Success comes from hard work, perseverance, and faith in God.'
}


def set_font_style(style, font_name='Arial'):
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    style._element.rPr.rFonts.set(qn('w:hAnsi'), font_name)
    style._element.rPr.rFonts.set(qn('w:cs'), font_name)


def add_section(doc, title, content=None, bullet_items=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)
    run.font.name = 'Arial'

    if content:
        paragraph.add_run('\n' + content)
    if bullet_items:
        for item in bullet_items:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            run = p.add_run(f'• {item}')
            run.font.size = Pt(11)
            run.font.name = 'Arial'


def add_contact_line(doc, label, value):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'{label}: {value}')
    run.font.size = Pt(10)
    run.font.name = 'Arial'


doc = Document()
styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal'].font.size = Pt(11)
set_font_style(styles['Normal'])
for style_name in ['Heading 1', 'Heading 2', 'Heading 3', 'Title', 'Subtitle']:
    if style_name in styles:
        set_font_style(styles[style_name])

section = doc.sections[0]
section.left_margin = Inches(0.6)
section.right_margin = Inches(0.6)
section.top_margin = Inches(0.4)
section.bottom_margin = Inches(0.4)

header = doc.add_paragraph()
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
header_run = header.add_run(profile_data['name'])
header_run.font.size = Pt(22)
header_run.font.bold = True
header_run.font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)
header_run.font.name = 'Arial'

role = doc.add_paragraph(profile_data['role'])
role.alignment = WD_ALIGN_PARAGRAPH.CENTER
role.runs[0].font.name = 'Arial'
role.runs[0].font.size = Pt(12)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.add_run(f'{profile_data["address"]} | {profile_data["email"]} | {profile_data["mobile"]} | {profile_data["facebook"]}').font.name = 'Arial'
contact.runs[0].font.size = Pt(10)

doc.add_paragraph()

# Photo and summary
photo_paragraph = doc.add_paragraph()
photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
if image_path.exists():
    photo_paragraph.add_run().add_picture(str(image_path), width=Inches(1.5))

summary_paragraph = doc.add_paragraph()
summary_paragraph.add_run('Professional Summary').bold = True
summary_paragraph.add_run('\n' + profile_data['summary'])
summary_paragraph.runs[0].font.name = 'Arial'
summary_paragraph.runs[0].font.size = Pt(13)
summary_paragraph.runs[0].font.color.rgb = RGBColor(0x0F, 0x4C, 0x81)
summary_paragraph.runs[1].font.name = 'Arial'
summary_paragraph.runs[1].font.size = Pt(11)

add_section(doc, 'Career Goals', profile_data['career_goals'])
add_section(doc, 'Education', bullet_items=profile_data['education'])
add_section(doc, 'Skills', content='Technical Skills: ' + ', '.join(profile_data['technical_skills']) + '\nProgramming Languages: ' + ', '.join(profile_data['programming_languages']) + '\nSoftware Applications: ' + ', '.join(profile_data['software_applications']) + '\nSoft Skills: ' + ', '.join(profile_data['soft_skills']))
add_section(doc, 'Featured Project', content=f'{profile_data["project_title"]}\n{profile_data["project_description"]}\nLanguage: {profile_data["project_language"]}\nStatus: {profile_data["project_status"]}\nLink: {profile_data["project_link"]}')
add_section(doc, 'Achievements & Certifications', bullet_items=profile_data['achievements'])
add_section(doc, 'Personal Motto', profile_data['motto'])

doc.save(output_path)
print(f'Resume created: {output_path}')
